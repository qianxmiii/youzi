"""从 Excel 指定列批量下载图片链接并打成 zip（支持进度任务）。"""

from __future__ import annotations

import io
import re
import threading
import time
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable
from urllib.parse import unquote, urlparse

import requests

_URL_RE = re.compile(r"^https?://", re.IGNORECASE)
_SAFE_NAME_RE = re.compile(r"[^\w.\-()+\[\]\u4e00-\u9fff]+", re.UNICODE)
_MAX_URLS = 500
_MAX_WORKERS = 16
_CONNECT_TIMEOUT = 10
_READ_TIMEOUT = 60
_MAX_BYTES = 20 * 1024 * 1024
_JOB_TTL_SEC = 3600

ProgressCb = Callable[[dict[str, Any]], None]


def _col_letter_to_index(letter: str) -> int:
    text = (letter or "P").strip().upper()
    if not text or not text.isalpha():
        raise ValueError("列字母无效，例如 P")
    n = 0
    for ch in text:
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n - 1


def _safe_filename(name: str, fallback: str) -> str:
    base = (name or "").strip() or fallback
    base = Path(base).name
    base = _SAFE_NAME_RE.sub("_", base).strip("._") or fallback
    return base[:180]


def _filename_from_url(url: str, index: int) -> str:
    path = unquote(urlparse(url).path or "")
    name = Path(path).name
    if not name or "." not in name:
        return f"image_{index:04d}.bin"
    return _safe_filename(name, f"image_{index:04d}.bin")


def _extract_urls_from_xls(data: bytes, *, col_index: int) -> list[str]:
    import xlrd

    wb = xlrd.open_workbook(file_contents=data)
    sh = wb.sheet_by_index(0)
    if col_index < 0 or col_index >= sh.ncols:
        raise ValueError(f"列索引超出范围（共 {sh.ncols} 列）")
    urls: list[str] = []
    for r in range(1, sh.nrows):
        raw = sh.cell_value(r, col_index)
        text = str(raw or "").strip()
        if _URL_RE.match(text):
            urls.append(text)
    return urls


def _extract_urls_from_xlsx(data: bytes, *, col_index: int) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sh = wb.active
        urls: list[str] = []
        for i, row in enumerate(
            sh.iter_rows(min_col=col_index + 1, max_col=col_index + 1, values_only=True)
        ):
            if i == 0:
                continue
            text = str(row[0] or "").strip() if row else ""
            if _URL_RE.match(text):
                urls.append(text)
        return urls
    finally:
        wb.close()


def extract_image_urls(data: bytes, *, filename: str, column: str = "P") -> list[str]:
    col_index = _col_letter_to_index(column)
    ext = Path(filename or "").suffix.lower()
    if ext == ".xls":
        urls = _extract_urls_from_xls(data, col_index=col_index)
    elif ext in (".xlsx", ".xlsm"):
        urls = _extract_urls_from_xlsx(data, col_index=col_index)
    else:
        try:
            urls = _extract_urls_from_xlsx(data, col_index=col_index)
        except Exception:
            urls = _extract_urls_from_xls(data, col_index=col_index)
    seen: set[str] = set()
    unique: list[str] = []
    for u in urls:
        if u not in seen:
            seen.add(u)
            unique.append(u)
    if len(unique) > _MAX_URLS:
        raise ValueError(f"链接数量 {len(unique)} 超过上限 {_MAX_URLS}")
    if not unique:
        raise ValueError(f"未在 {column} 列找到有效的 http(s) 图片链接")
    return unique


def _download_one(url: str, index: int) -> tuple[str, bytes | None, str | None]:
    name = _filename_from_url(url, index)
    try:
        resp = requests.get(
            url,
            timeout=(_CONNECT_TIMEOUT, _READ_TIMEOUT),
            stream=True,
        )
        resp.raise_for_status()
        chunks: list[bytes] = []
        total = 0
        for chunk in resp.iter_content(chunk_size=256 * 1024):
            if not chunk:
                continue
            total += len(chunk)
            if total > _MAX_BYTES:
                return name, None, f"超过单文件上限 {_MAX_BYTES // (1024 * 1024)}MB"
            chunks.append(chunk)
        if not chunks:
            return name, None, "空响应"
        if "." not in name:
            ctype = (resp.headers.get("Content-Type") or "").split(";")[0].strip().lower()
            ext_map = {
                "image/png": ".png",
                "image/jpeg": ".jpg",
                "image/jpg": ".jpg",
                "image/gif": ".gif",
                "image/webp": ".webp",
            }
            name = _safe_filename(name + ext_map.get(ctype, ".bin"), f"image_{index:04d}.bin")
        return name, b"".join(chunks), None
    except Exception as exc:  # noqa: BLE001
        return name, None, str(exc) or "下载失败"


def build_image_zip_from_urls(
    urls: list[str],
    *,
    on_progress: ProgressCb | None = None,
) -> tuple[bytes, dict[str, Any]]:
    """并发下载并打包；on_progress 每次完成一张回调。"""
    results: list[tuple[int, str, bytes | None, str | None]] = []
    done = 0
    failed_count = 0
    total = len(urls)

    def _emit() -> None:
        if on_progress:
            on_progress(
                {
                    "done": done,
                    "total": total,
                    "failed": failed_count,
                    "status": "running",
                }
            )

    with ThreadPoolExecutor(max_workers=_MAX_WORKERS) as pool:
        futures = {
            pool.submit(_download_one, url, i + 1): (i, url) for i, url in enumerate(urls)
        }
        for fut in as_completed(futures):
            i, _url = futures[fut]
            name, content, err = fut.result()
            results.append((i, name, content, err))
            done += 1
            if content is None:
                failed_count += 1
            _emit()

    results.sort(key=lambda x: x[0])
    used_names: dict[str, int] = {}
    ok = 0
    failed: list[dict[str, str]] = []
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for i, name, content, err in results:
            if content is None:
                failed.append({"url": urls[i], "name": name, "error": err or "下载失败"})
                continue
            base = name
            if base in used_names:
                used_names[base] += 1
                stem = Path(base).stem
                suffix = Path(base).suffix
                final = f"{stem}_{used_names[base]}{suffix}"
            else:
                used_names[base] = 1
                final = base
            zf.writestr(final, content)
            ok += 1

    if ok == 0:
        raise ValueError("全部图片下载失败，未生成压缩包")

    return buf.getvalue(), {
        "total": total,
        "downloaded": ok,
        "failed": len(failed),
        "errors": failed[:20],
    }


def build_image_zip_from_excel(
    data: bytes,
    *,
    filename: str,
    column: str = "P",
    on_progress: ProgressCb | None = None,
) -> tuple[bytes, dict[str, Any]]:
    urls = extract_image_urls(data, filename=filename, column=column)
    if on_progress:
        on_progress({"done": 0, "total": len(urls), "failed": 0, "status": "running"})
    zip_bytes, stats = build_image_zip_from_urls(urls, on_progress=on_progress)
    stats["column"] = column.upper()
    return zip_bytes, stats


@dataclass
class ImageZipJob:
    id: str
    status: str = "pending"  # pending | running | done | failed
    total: int = 0
    done: int = 0
    failed: int = 0
    downloaded: int = 0
    error: str | None = None
    column: str = "P"
    zip_bytes: bytes | None = None
    created_at: float = field(default_factory=time.time)
    updated_at: float = field(default_factory=time.time)

    def to_api(self) -> dict[str, Any]:
        return {
            "jobId": self.id,
            "status": self.status,
            "total": self.total,
            "done": self.done,
            "failed": self.failed,
            "downloaded": self.downloaded,
            "error": self.error,
            "column": self.column,
            "ready": self.status == "done" and self.zip_bytes is not None,
        }


class ImageZipJobStore:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._jobs: dict[str, ImageZipJob] = {}

    def _purge_locked(self) -> None:
        now = time.time()
        dead = [jid for jid, j in self._jobs.items() if now - j.updated_at > _JOB_TTL_SEC]
        for jid in dead:
            self._jobs.pop(jid, None)

    def create(
        self,
        data: bytes,
        *,
        filename: str,
        column: str = "P",
    ) -> ImageZipJob:
        urls = extract_image_urls(data, filename=filename, column=column)
        job = ImageZipJob(
            id=str(uuid.uuid4()),
            status="pending",
            total=len(urls),
            column=column.upper(),
        )
        with self._lock:
            self._purge_locked()
            self._jobs[job.id] = job

        def _run() -> None:
            with self._lock:
                job.status = "running"
                job.updated_at = time.time()

            def on_progress(info: dict[str, Any]) -> None:
                with self._lock:
                    job.done = int(info.get("done") or 0)
                    job.total = int(info.get("total") or job.total)
                    job.failed = int(info.get("failed") or 0)
                    job.updated_at = time.time()

            try:
                zip_bytes, stats = build_image_zip_from_urls(urls, on_progress=on_progress)
                with self._lock:
                    job.zip_bytes = zip_bytes
                    job.downloaded = int(stats.get("downloaded") or 0)
                    job.failed = int(stats.get("failed") or 0)
                    job.done = job.total
                    job.status = "done"
                    job.updated_at = time.time()
            except Exception as exc:  # noqa: BLE001
                with self._lock:
                    job.status = "failed"
                    job.error = str(exc) or "打包失败"
                    job.updated_at = time.time()

        threading.Thread(target=_run, name=f"image-zip-{job.id[:8]}", daemon=True).start()
        return job

    def get(self, job_id: str) -> ImageZipJob | None:
        with self._lock:
            self._purge_locked()
            return self._jobs.get(job_id)


image_zip_jobs = ImageZipJobStore()
