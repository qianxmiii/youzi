"""运单 Excel 导出：承运商中文名、转单号、品名。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from youzi_v2.services.shipment_excel import build_export_excel_bytes


def test_export_uses_carrier_name_zh_and_new_columns() -> None:
    data = build_export_excel_bytes(
        [
            {
                "shipmentNo": "DPSECO1",
                "statusCode": "IN_TRANSIT",
                "customerNo": "FBA111",
                "refShipmentNo": "PO-1",
                "carrierCode": "1724258253196189697",
                "carrierNameZh": "天图FBA",
                "trackingNumber": "1Z999",
                "productName": "相机浮力带",
                "ctns": 2,
                "etd": "2026-05-01 10:00:00",
                "atd": "2026-05-02 12:00:00",
                "eta": "2026-05-20 08:00:00",
                "ata": "2026-05-21 09:30:00",
            }
        ]
    )
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "承运商" in headers
    assert "转单号" in headers
    assert "品名" in headers
    assert headers[-2:] == ["内部轨迹时间", "内部最新轨迹"]
    assert headers.index("品名") == headers.index("供应商") + 1
    assert headers[headers.index("ETD") : headers.index("ATA") + 1] == [
        "ETD",
        "ATD",
        "ETA",
        "ATA",
    ]
    row = {headers[i]: ws[2][i].value for i in range(len(headers))}
    assert row["承运商"] == "天图FBA"
    assert row["转单号"] == "1Z999"
    assert row["品名"] == "相机浮力带"
    assert row["运单号"] == "DPSECO1"
    assert row["ETD"] == "2026-05-01"
    assert row["ATD"] == "2026-05-02"
    assert row["ETA"] == "2026-05-20"
    assert row["ATA"] == "2026-05-21"


def test_export_carrier_falls_back_to_code() -> None:
    data = build_export_excel_bytes(
        [
            {
                "shipmentNo": "DPSECO2",
                "carrierCode": "1724258253196189697",
                "carrierNameZh": None,
            }
        ]
    )
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    row = {headers[i]: ws[2][i].value for i in range(len(headers))}
    assert row["承运商"] == "1724258253196189697"
