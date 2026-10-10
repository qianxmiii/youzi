"""按运单 Excel 映射生成文件重命名目标名。"""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "shipmentNo": ("运单号", "shipment_no", "shipmentno", "odd"),
    "country": ("国家", "国家代码", "country", "country_code", "countrycode"),
    "addressCode": ("仓库代码", "派送仓库", "仓库", "address_code", "addresscode", "warehouse"),
    "ctns": ("件数", "箱数", "ctns", "cartons"),
    # 可选：相同文件夹名归入同一子目录
    "folderName": ("文件夹名", "目录名", "folder", "folder_name", "foldername", "子文件夹"),
}

_SHEET_PREFERENCE = ("运单信息", "Sheet1", "sheet1")

_INVALID_FS = re.compile(r'[\\/:*?"<>|]+')


def _norm_header(value: Any) -> str:
    text = str(value or "").replace("\r", "\n")
    text = " ".join(part.strip() for part in text.split("\n") if part.strip())
    return text.strip().lower().replace(" ", "").replace("_", "")


def _header_matches(cell_text: str, alias: str) -> bool:
    """精确或前缀匹配表头（兼容换行/括号说明）。"""
    wanted = _norm_header(alias)
    if not wanted:
        return False
    if cell_text == wanted:
        return True
    if cell_text.startswith(wanted):
        rest = cell_text[len(wanted) :]
        return not rest or rest[0] in "（("
    return False


def _map_headers(row: tuple[Any, ...]) -> dict[str, int]:
    found: dict[str, int] = {}
    for idx, cell in enumerate(row):
        text = _norm_header(cell)
        if not text:
            continue
        for field, aliases in _HEADER_ALIASES.items():
            if field in found:
                continue
            for alias in aliases:
                if _header_matches(text, alias):
                    found[field] = idx
                    break
    return found


def sanitize_filename_part(value: Any) -> str:
    text = str(value if value is not None else "").strip()
    text = _INVALID_FS.sub("-", text)
    text = re.sub(r"\s+", "", text)
    return text


def build_rename_target(
    *,
    shipment_no: str,
    country: str,
    address_code: str,
    ctns: str,
) -> str:
    return "-".join(
        [
            sanitize_filename_part(shipment_no),
            sanitize_filename_part(country),
            sanitize_filename_part(address_code),
            sanitize_filename_part(ctns),
        ]
    )


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_from_xlsx(content: bytes) -> list[list[tuple[Any, ...]]]:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        ordered: list[str] = []
        for name in _SHEET_PREFERENCE:
            if name in wb.sheetnames and name not in ordered:
                ordered.append(name)
        for name in wb.sheetnames:
            if name not in ordered:
                ordered.append(name)

        sheets: list[list[tuple[Any, ...]]] = []
        for name in ordered:
            ws = wb[name]
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                sheets.append(rows)
        return sheets
    finally:
        wb.close()


def _rows_from_xls(content: bytes) -> list[list[tuple[Any, ...]]]:
    import xlrd

    wb = xlrd.open_workbook(file_contents=content)
    sheets: list[list[tuple[Any, ...]]] = []
    preferred = [wb.sheet_by_name(name) for name in _SHEET_PREFERENCE if name in wb.sheet_names()]
    others = [wb.sheet_by_index(i) for i in range(wb.nsheets)]
    seen: set[str] = set()
    ordered = []
    for sh in preferred + others:
        if sh.name in seen:
            continue
        seen.add(sh.name)
        ordered.append(sh)

    for sh in ordered:
        rows: list[tuple[Any, ...]] = []
        for r in range(sh.nrows):
            rows.append(tuple(sh.row_values(r)))
        if rows:
            sheets.append(rows)
    return sheets


def _load_workbook_rows(content: bytes, *, filename: str = "") -> list[list[tuple[Any, ...]]]:
    ext = Path(filename or "").suffix.lower()
    if ext == ".xls":
        return _rows_from_xls(content)
    if ext in {".xlsx", ".xlsm"}:
        return _rows_from_xlsx(content)
    try:
        return _rows_from_xlsx(content)
    except Exception:
        return _rows_from_xls(content)


def _find_header(sheets: Iterable[list[tuple[Any, ...]]]) -> tuple[int, dict[str, int], list[tuple[Any, ...]]]:
    for rows in sheets:
        for i, row in enumerate(rows[:30]):
            mapped = _map_headers(tuple(row or ()))
            if (
                "shipmentNo" in mapped
                and "country" in mapped
                and "addressCode" in mapped
                and "ctns" in mapped
            ):
                return i, mapped, rows
    return -1, {}, []


def parse_shipment_rename_excel(content: bytes, *, filename: str = "") -> dict[str, Any]:
    """解析入仓/运单明细 Excel，返回重命名映射。

    可选列「文件夹名」：有值时相同文件夹名归入同一子目录。
    """
    sheets = _load_workbook_rows(content, filename=filename)
    if not sheets:
        raise ValueError("Excel 为空")

    header_idx, col_map, rows = _find_header(sheets)
    if header_idx < 0:
        raise ValueError("未找到表头，需包含：运单号、国家、仓库代码、件数")

    items: list[dict[str, str]] = []
    errors: list[str] = []
    seen: dict[str, int] = {}
    for row_no, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        cells = tuple(row or ())
        shipment_no = _cell_str(cells[col_map["shipmentNo"]] if col_map["shipmentNo"] < len(cells) else "")
        if not shipment_no:
            continue
        country = _cell_str(cells[col_map["country"]] if col_map["country"] < len(cells) else "")
        address_code = _cell_str(
            cells[col_map["addressCode"]] if col_map["addressCode"] < len(cells) else ""
        )
        ctns = _cell_str(cells[col_map["ctns"]] if col_map["ctns"] < len(cells) else "")
        folder_raw = ""
        if "folderName" in col_map and col_map["folderName"] < len(cells):
            folder_raw = _cell_str(cells[col_map["folderName"]])
        folder_name = sanitize_filename_part(folder_raw) if folder_raw else ""
        if not country or not address_code or not ctns:
            errors.append(f"第 {row_no} 行（{shipment_no}）国家/仓库代码/件数不完整，已跳过")
            continue
        key = shipment_no.upper()
        if key in seen:
            errors.append(f"运单号重复：{shipment_no}（第 {seen[key]} 行与第 {row_no} 行），保留首次")
            continue
        seen[key] = row_no
        target_stem = build_rename_target(
            shipment_no=shipment_no,
            country=country,
            address_code=address_code,
            ctns=ctns,
        )
        items.append(
            {
                "shipmentNo": shipment_no,
                "country": country,
                "addressCode": address_code,
                "ctns": ctns,
                "folderName": folder_name,
                "targetStem": target_stem,
            }
        )

    if not items:
        raise ValueError("未解析到有效运单行（请检查国家、仓库代码、件数是否有值）")
    return {"items": items, "total": len(items), "errors": errors}
