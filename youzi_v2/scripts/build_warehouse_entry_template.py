"""生成进仓单参考模板 youzi_v2/config/templates/warehouse_entry_template.xlsx"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "config" / "templates" / "warehouse_entry_template.xlsx"


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "进仓单"

    widths = {"A": 6, "B": 14, "C": 12, "D": 10, "E": 14, "F": 12, "G": 18, "H": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    thin = Side(style="thin", color="999999")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="E8F4FC")
    label_fill = PatternFill("solid", fgColor="F5F5F5")

    title_font = Font(name="Microsoft YaHei", size=16, bold=True)
    label_font = Font(name="Microsoft YaHei", size=11, bold=True)
    body_font = Font(name="Microsoft YaHei", size=11)
    placeholder_font = Font(name="Microsoft YaHei", size=11, color="C2410C")
    table_header_font = Font(name="Microsoft YaHei", size=11, bold=True)
    hint_font = Font(name="Microsoft YaHei", size=9, color="666666")

    def merge_title(r1: int, c1: int, r2: int, c2: int, value: str) -> None:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        cell = ws.cell(row=r1, column=c1, value=value)
        cell.font = title_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    merge_title(1, 1, 1, 8, "送 货 进 仓 单")

    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 22
    meta_rows = [
        (2, "进仓日期", "{日期}", "合计票数", "{票数}"),
        (3, "收货仓", "{收货仓}", "制单人", "{制单人}"),
    ]
    for row, l1, v1, l2, v2 in meta_rows:
        ws.cell(row=row, column=1, value=l1).font = label_font
        ws.cell(row=row, column=1).fill = label_fill
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        c1 = ws.cell(row=row, column=2, value=v1)
        c1.font = placeholder_font
        c1.alignment = Alignment(vertical="center")
        ws.cell(row=row, column=4, value=l2).font = label_font
        ws.cell(row=row, column=4).fill = label_fill
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        c2 = ws.cell(row=row, column=5, value=v2)
        c2.font = placeholder_font if "{" in v2 else body_font
        c2.alignment = Alignment(vertical="center")

    ws.cell(row=4, column=1, value="收货地址").font = label_font
    ws.cell(row=4, column=1).fill = label_fill
    ws.cell(row=4, column=1).alignment = Alignment(vertical="top")
    ws.merge_cells(start_row=4, start_column=2, end_row=7, end_column=8)
    addr = ws.cell(row=4, column=2, value="{收货地址}")
    addr.font = placeholder_font
    addr.alignment = Alignment(vertical="top", wrap_text=True)
    for r in range(4, 8):
        ws.row_dimensions[r].height = 20

    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=8)
    hint = ws.cell(
        row=8,
        column=1,
        value="说明：下方明细可手工填，或由系统/DPS 从第 10 行表头起粘贴。橙色 {…} 为占位符，接入程序后替换。",
    )
    hint.font = hint_font
    hint.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[8].height = 28
    ws.row_dimensions[9].height = 8

    headers = ["序号", "运单号", "客户", "件数", "渠道", "客户单号", "货件号", "备注"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=10, column=i, value=h)
        cell.font = table_header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_all
    ws.row_dimensions[10].height = 24

    example = ["1", "{运单号}", "{客户}", "{件数}", "{渠道}", "{客户单号}", "{货件号}", ""]
    for i, val in enumerate(example, start=1):
        cell = ws.cell(row=11, column=i, value=val)
        cell.font = placeholder_font if val.startswith("{") else body_font
        cell.border = border_all
        cell.alignment = Alignment(
            horizontal="center" if i in (1, 4) else "left",
            vertical="center",
        )

    for r in range(12, 16):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c, value=r - 10 if c == 1 else "")
            cell.border = border_all
            cell.alignment = Alignment(vertical="center", horizontal="center" if c == 1 else "left")

    leg = wb.create_sheet("占位符说明")
    leg.column_dimensions["A"].width = 22
    leg.column_dimensions["B"].width = 52
    leg.append(["占位符", "含义 / 数据来源建议"])
    for cell in leg[1]:
        cell.font = Font(bold=True)
    for a, b in [
        ("{日期}", "进仓单日期，如 2026-10-18"),
        ("{票数}", "勾选运单数量 N"),
        ("{收货仓}", "深圳仓 / 义乌仓（originWarehouseCode）"),
        ("{收货地址}", "config.json → dps_warehouse.SHENZHEN 或 YIWU"),
        ("{制单人}", "可选：登录用户或留空"),
        ("{运单号}", "shipmentNo"),
        ("{客户}", "customer"),
        ("{件数}", "ctns"),
        ("{渠道}", "channelNameZh / channelCode"),
        ("{客户单号}", "customerNo"),
        ("{货件号}", "customerShipmentId"),
    ]:
        leg.append([a, b])

    wb.save(OUT)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    main()
