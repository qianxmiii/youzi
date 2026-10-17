"""按 config/templates/warehouse_entry_template.xlsx 生成进仓单。"""

from __future__ import annotations

import io
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet

_FONT_YHUI_BOLD = Font(name="Microsoft YaHei UI", bold=True)
_THIN_SIDE = Side(style="thin", color="000000")
_BORDER_ALL = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)

from .dps_receive_warehouse import receive_warehouse_label_zh
from .shipment_query_config import load_config_json

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE_PATH = ROOT / "config" / "templates" / "warehouse_entry_template.xlsx"
DATA_SHEET_NAME = "进仓单"
LEGEND_SHEET_NAME = "占位符说明"

_PLACEHOLDER_RE = re.compile(r"\{([^}]+)\}")

FOOTER_REMARKS_ROWS = 2
FOOTER_ADDRESS_ROWS = 3
DETAIL_ROW_HEIGHT = 25
DEFAULT_REMARKS_LABEL = "特殊备注栏:\n(特殊要求）"
_EMPTY_WAREHOUSE_DISPLAY = "——"


def _footer_compact_row_count() -> int:
    return 1 + FOOTER_REMARKS_ROWS + FOOTER_ADDRESS_ROWS


def resolve_warehouse_entry_template_path(config_path: Path | None) -> Path:
    if config_path and config_path.is_file():
        config = load_config_json(config_path)
        raw = config.get("warehouse_entry_template")
        if isinstance(raw, str) and raw.strip():
            p = Path(raw.strip())
            if not p.is_absolute():
                repo_root = config_path.resolve().parent.parent
                p = repo_root / p
            if p.is_file():
                return p
    if DEFAULT_TEMPLATE_PATH.is_file():
        return DEFAULT_TEMPLATE_PATH
    raise FileNotFoundError(
        f"进仓单模板不存在：{DEFAULT_TEMPLATE_PATH}（或在 config.json 配置 warehouse_entry_template）"
    )


_COUNTRY_CODE_TO_ZH: dict[str, str] = {
    "US": "美国",
    "CA": "加拿大",
    "GB": "英国",
    "DE": "德国",
    "FR": "法国",
    "EU": "欧洲",
    "AU": "澳大利亚",
    "JP": "日本",
}


def _country_label(row: dict[str, Any]) -> str:
    code = str(row.get("countryCode") or "").strip().upper()
    if not code:
        return ""
    return _COUNTRY_CODE_TO_ZH.get(code, code)


def _channel_label(row: dict[str, Any]) -> str:
    zh = str(row.get("channelNameZh") or "").strip()
    if zh:
        return zh
    return str(row.get("channelCode") or "").strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _delivery_warehouse_display(value: Any) -> str:
    text = _cell_text(value)
    return text if text else _EMPTY_WAREHOUSE_DISPLAY


def _find_row_containing(ws: Worksheet, *needles: str) -> int | None:
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            text = _cell_text(ws.cell(r, c).value)
            if not text:
                continue
            for n in needles:
                if n in text:
                    return r
    return None


def _find_detail_template_row(ws: Worksheet, header_row: int | None = None) -> int:
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if _cell_text(ws.cell(r, c).value) == "{运单号}":
                return r
    if header_row is not None and header_row + 1 <= ws.max_row:
        return header_row + 1
    raise ValueError("模板中未找到明细行（首行明细保留 {运单号}，或紧挨表头下一行）")


def _find_table_header_row(ws: Worksheet) -> int | None:
    for r in range(1, ws.max_row + 1):
        headers = {
            _cell_text(ws.cell(r, c).value)
            for c in range(1, ws.max_column + 1)
        }
        if "运单号" in headers and "件数" in headers:
            return r
    return None


def _find_total_row(ws: Worksheet) -> int:
    """合计行：优先 {合计件数}，其次首行仅含「合计」的单元格（避免误匹配表头）。"""
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            text = _cell_text(ws.cell(r, c).value)
            if text == "{合计件数}":
                return r
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            text = _cell_text(ws.cell(r, c).value)
            if "合计" in text and "件数" not in text.replace("合计件数", ""):
                if text.startswith("合计") or text in ("合计", "合计：", "合计:"):
                    return r
    raise ValueError("模板中未找到合计行（需 {合计件数} 或「合计：」）")


def _unmerge_intersecting_rows(ws: Worksheet, row_start: int, row_end: int) -> None:
    """仅解除与明细行区间相交的合并（删行/插行后避免 unmerge_cells KeyError）。"""
    if row_end < row_start:
        return
    for merged in list(ws.merged_cells.ranges):
        if merged.max_row < row_start or merged.min_row > row_end:
            continue
        ws.merged_cells.remove(merged)


def _drop_merges_in_block(ws: Worksheet, row_start: int, row_end: int, max_col: int) -> None:
    for merged in list(ws.merged_cells.ranges):
        if merged.max_row < row_start or merged.min_row > row_end:
            continue
        if merged.min_col > max_col:
            continue
        ws.merged_cells.remove(merged)


def load_warehouse_entry_special_remarks(
    config_path: Path | None,
    warehouse_key: str,
) -> str:
    """config.json → dps_warehouse_entry_remarks（按仓键或全局字符串）。"""
    if not config_path or not config_path.is_file():
        return ""
    config = load_config_json(config_path)
    raw = config.get("dps_warehouse_entry_remarks")
    key = warehouse_key.strip().upper()
    if isinstance(raw, dict):
        for k, v in raw.items():
            if str(k).strip().upper() == key and isinstance(v, str):
                return v.strip()
        return ""
    if isinstance(raw, str):
        return raw.strip()
    return ""


def _trim_spare_detail_rows(ws: Worksheet, detail_row: int, total_row: int) -> int:
    """去掉模板里序号 2、3… 的空白明细行，只保留首行样式。"""
    if total_row <= detail_row + 1:
        return total_row
    remove = total_row - detail_row - 1
    ws.delete_rows(detail_row + 1, remove)
    return total_row - remove


def _expand_detail_rows(
    ws: Worksheet,
    detail_row: int,
    total_row: int,
    need: int,
    max_col: int,
) -> int:
    """在合计行上方扩到 need 行明细，返回新的合计行号。"""
    if need <= 1:
        return total_row
    extra = need - 1
    ws.insert_rows(total_row, amount=extra)
    for offset in range(extra):
        _copy_row_style(ws, detail_row, total_row + offset, max_col)
    return total_row + extra


def _snapshot_footer_rows(
    ws: Worksheet,
    total_row: int,
    max_col: int,
) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for r in range(total_row, ws.max_row + 1):
        rows.append([ws.cell(r, c).value for c in range(1, max_col + 1)])
    return rows


def _extract_remarks_label(footer_rows: list[list[Any]]) -> str:
    for row in footer_rows:
        for val in row:
            if isinstance(val, str) and "特殊备注" in val:
                return val.strip()
    return DEFAULT_REMARKS_LABEL


def _apply_grid_borders(
    ws: Worksheet,
    row_start: int,
    row_end: int,
    col_start: int,
    col_end: int,
) -> None:
    for r in range(row_start, row_end + 1):
        for c in range(col_start, col_end + 1):
            ws.cell(r, c).border = _BORDER_ALL


def _style_footer_text_cell(
    cell,
    *,
    horizontal: str = "center",
    vertical: str = "center",
    wrap: bool = False,
) -> None:
    cell.font = _FONT_YHUI_BOLD
    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical=vertical,
        wrap_text=wrap,
    )
    cell.border = _BORDER_ALL


def _restore_footer_merges(
    ws: Worksheet,
    total_row: int,
    max_col: int,
    remarks_rows: int,
    address_rows: int,
) -> None:
    remarks_row = total_row + 1
    address_row = total_row + 1 + remarks_rows
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=max_col - 1)
    ws.merge_cells(
        start_row=remarks_row,
        start_column=1,
        end_row=remarks_row + remarks_rows - 1,
        end_column=1,
    )
    ws.merge_cells(
        start_row=remarks_row,
        start_column=2,
        end_row=remarks_row + remarks_rows - 1,
        end_column=max_col,
    )
    ws.merge_cells(
        start_row=address_row,
        start_column=1,
        end_row=address_row + address_rows - 1,
        end_column=1,
    )
    ws.merge_cells(
        start_row=address_row,
        start_column=2,
        end_row=address_row + address_rows - 1,
        end_column=max_col,
    )


def _write_compact_footer(
    ws: Worksheet,
    start_row: int,
    max_col: int,
    *,
    total_ctns: int,
    receive_address: str,
    special_remarks: str,
    remarks_label: str,
) -> None:
    """紧凑页脚：合计 1 行 + 备注 2 行 + 地址 3 行（不复制模板空行）。"""
    remarks_rows = FOOTER_REMARKS_ROWS
    address_rows = FOOTER_ADDRESS_ROWS
    block_end = start_row + _footer_compact_row_count() - 1
    _drop_merges_in_block(ws, start_row, block_end, max_col)

    for r in range(start_row, block_end + 1):
        for c in range(1, max_col + 1):
            ws.cell(r, c, value=None)

    total_r = start_row
    remarks_r = start_row + 1
    address_r = start_row + 1 + remarks_rows

    _apply_grid_borders(ws, total_r, block_end, 1, max_col)

    ws.cell(total_r, 1, value="合计：")
    ws.cell(total_r, max_col, value=str(total_ctns))

    label = remarks_label.strip() or DEFAULT_REMARKS_LABEL
    ws.cell(remarks_r, 1, value=label)
    if special_remarks.strip():
        ws.cell(remarks_r, 2, value=special_remarks.strip())

    ws.cell(address_r, 1, value="收货地址")
    ws.cell(address_r, 2, value=receive_address.strip())

    ws.row_dimensions[total_r].height = DETAIL_ROW_HEIGHT
    ws.row_dimensions[remarks_r].height = 36
    ws.row_dimensions[address_r].height = 66

    _restore_footer_merges(ws, total_r, max_col, remarks_rows, address_rows)
    _apply_footer_text_styles(
        ws,
        total_r,
        max_col,
        remarks_rows=remarks_rows,
        special_remarks=special_remarks,
        total_ctns=total_ctns,
    )


def _apply_footer_text_styles(
    ws: Worksheet,
    total_r: int,
    max_col: int,
    *,
    remarks_rows: int,
    special_remarks: str,
    total_ctns: int,
) -> None:
    remarks_r = total_r + 1
    address_r = total_r + 1 + remarks_rows
    ws.cell(total_r, max_col, value=str(total_ctns))
    _style_footer_text_cell(ws.cell(total_r, 1), horizontal="right", vertical="center")
    _style_footer_text_cell(ws.cell(total_r, max_col), horizontal="center", vertical="center")
    _style_footer_text_cell(
        ws.cell(remarks_r, 1),
        horizontal="center",
        vertical="center",
        wrap=True,
    )
    if special_remarks.strip():
        _style_footer_text_cell(
            ws.cell(remarks_r, 2),
            horizontal="center",
            vertical="center",
            wrap=True,
        )
    _style_footer_text_cell(
        ws.cell(address_r, 1),
        horizontal="center",
        vertical="center",
        wrap=True,
    )
    _style_footer_text_cell(
        ws.cell(address_r, 2),
        horizontal="center",
        vertical="center",
        wrap=True,
    )


def _copy_row_style(ws: Worksheet, src_row: int, dst_row: int, max_col: int) -> None:
    for c in range(1, max_col + 1):
        src = ws.cell(src_row, c)
        dst = ws.cell(dst_row, c)
        dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)


def _replace_placeholders_in_cell(value: Any, mapping: dict[str, str]) -> Any:
    if not isinstance(value, str) or "{" not in value:
        return value

    def repl(match: re.Match[str]) -> str:
        key = match.group(1).strip()
        return mapping.get(key, match.group(0))

    return _PLACEHOLDER_RE.sub(repl, value)


def _apply_mapping_to_sheet(ws: Worksheet, mapping: dict[str, str]) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            new_val = _replace_placeholders_in_cell(cell.value, mapping)
            if new_val != cell.value:
                cell.value = new_val


WAREHOUSE_ENTRY_DETAIL_HEADERS: tuple[str, ...] = (
    "序号",
    "运单号",
    "客户单号",
    "目的国",
    "派送地址",
    "邮编",
    "件数",
)


def warehouse_entry_detail_table(
    shipments: list[dict[str, Any]],
) -> tuple[list[tuple[str, ...]], int]:
    """明细行（与模板表头一致）及件数合计。"""
    rows: list[tuple[str, ...]] = []
    total_ctns = 0
    for i, ship in enumerate(shipments, 1):
        fields = _row_field_mapping(ship)
        rows.append(
            (
                str(i),
                fields["运单号"],
                fields["客户单号"],
                fields["目的国"],
                fields["派送地址"],
                fields["邮编"],
                fields["件数"],
            )
        )
        try:
            total_ctns += int(ship.get("ctns") or 0)
        except (TypeError, ValueError):
            pass
    return rows, total_ctns


def _row_field_mapping(shipment: dict[str, Any]) -> dict[str, str]:
    ctns = shipment.get("ctns")
    ctns_s = "" if ctns is None else _cell_text(ctns)
    return {
        "运单号": _cell_text(shipment.get("shipmentNo")),
        "客户单号": _cell_text(shipment.get("customerNo")),
        "渠道": _channel_label(shipment),
        "目的国": _country_label(shipment),
        "国家": _country_label(shipment),
        "派送仓库": _delivery_warehouse_display(shipment.get("addressCode")),
        "派送地址": _delivery_warehouse_display(shipment.get("addressCode")),
        "邮编": _cell_text(shipment.get("zipcode")),
        "件数": ctns_s,
        "备注": "",
    }


_HEADER_TO_FIELD: dict[str, str] = {
    "运单号": "运单号",
    "客户单号": "客户单号",
    "渠道": "渠道",
    "目的国": "目的国",
    "国家": "国家",
    "派送地址": "派送地址",
    "派送仓库": "派送仓库",
    "邮编": "邮编",
    "件数": "件数",
    "备注": "备注",
}


def _detail_column_map(ws: Worksheet, header_row: int, max_col: int) -> dict[int, str]:
    """列号 → 字段键（序号列用 __index__）。"""
    col_map: dict[int, str] = {}
    for c in range(1, max_col + 1):
        header = _cell_text(ws.cell(header_row, c).value)
        if header == "序号":
            col_map[c] = "__index__"
            continue
        field = _HEADER_TO_FIELD.get(header)
        if field:
            col_map[c] = field
    return col_map


def _write_shipment_to_row(
    ws: Worksheet,
    row_idx: int,
    col_map: dict[int, str],
    shipment: dict[str, Any],
    index: int,
) -> None:
    fields = _row_field_mapping(shipment)
    for col, key in col_map.items():
        if key == "__index__":
            ws.cell(row_idx, col, value=index)
        else:
            ws.cell(row_idx, col, value=fields.get(key, ""))


def build_warehouse_entry_xlsx_from_template(
    template_path: Path,
    shipments: list[dict[str, Any]],
    *,
    warehouse_key: str,
    warehouse_label: str,
    receive_address: str,
    special_remarks: str = "",
    when: datetime | None = None,
) -> bytes:
    if not shipments:
        raise ValueError("没有运单数据可写入进仓单")

    when = when or datetime.now()
    wb = load_workbook(template_path)
    if DATA_SHEET_NAME in wb.sheetnames:
        ws = wb[DATA_SHEET_NAME]
    else:
        ws = wb.active

    header_row = _find_table_header_row(ws)
    if header_row is None:
        raise ValueError("模板中未找到明细表头行（需含「运单号」「件数」等列名）")
    detail_row = _find_detail_template_row(ws, header_row=header_row)
    if detail_row <= header_row:
        detail_row = header_row + 1

    total_row = _find_total_row(ws)

    max_col = max(ws.max_column, 1)
    col_map = _detail_column_map(ws, header_row, max_col)
    if not any(v != "__index__" for v in col_map.values()):
        raise ValueError("模板表头行未识别到运单明细列（需含运单号、件数等）")

    footer_snapshot = _snapshot_footer_rows(ws, total_row, max_col)
    total_row = _trim_spare_detail_rows(ws, detail_row, total_row)
    need = len(shipments)
    total_row = _expand_detail_rows(ws, detail_row, total_row, need, max_col)
    _unmerge_intersecting_rows(ws, detail_row, detail_row + need - 1)

    new_total_row = detail_row + need
    if ws.max_row >= new_total_row:
        ws.delete_rows(new_total_row, ws.max_row - new_total_row + 1)
    footer_len = _footer_compact_row_count()
    remarks_label = _extract_remarks_label(footer_snapshot) if footer_snapshot else DEFAULT_REMARKS_LABEL
    if footer_snapshot:
        ws.insert_rows(new_total_row, amount=footer_len)

    total_ctns = 0
    for i, ship in enumerate(shipments):
        r = detail_row + i
        _write_shipment_to_row(ws, r, col_map, ship, i + 1)
        ws.row_dimensions[r].height = DETAIL_ROW_HEIGHT
        try:
            total_ctns += int(ship.get("ctns") or 0)
        except (TypeError, ValueError):
            pass

    header_map = {
        "日期": when.strftime("%Y-%m-%d"),
        "收货仓": warehouse_label or receive_warehouse_label_zh(warehouse_key),
        "票数": str(len(shipments)),
        "制单人": "",
    }
    _apply_mapping_to_sheet(ws, header_map)

    if footer_snapshot:
        _write_compact_footer(
            ws,
            new_total_row,
            max_col,
            total_ctns=total_ctns,
            receive_address=receive_address,
            special_remarks=special_remarks,
            remarks_label=remarks_label,
        )
        footer_end = new_total_row + footer_len - 1
        if ws.max_row > footer_end:
            ws.delete_rows(footer_end + 1, ws.max_row - footer_end)
    else:
        footer_end = detail_row + need - 1

    _apply_grid_borders(ws, header_row, footer_end, 1, max_col)
    if footer_snapshot:
        _apply_footer_text_styles(
            ws,
            new_total_row,
            max_col,
            remarks_rows=FOOTER_REMARKS_ROWS,
            special_remarks=special_remarks,
            total_ctns=total_ctns,
        )

    if LEGEND_SHEET_NAME in wb.sheetnames:
        del wb[LEGEND_SHEET_NAME]

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
