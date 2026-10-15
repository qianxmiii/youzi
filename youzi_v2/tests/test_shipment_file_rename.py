"""运单文件按 Excel 重命名：解析单测。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from youzi_v2.services.shipment_file_rename import (
    build_rename_target,
    parse_shipment_rename_excel,
)


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_build_rename_target():
    assert build_rename_target(
        shipment_no="DPSECO260812093",
        country="CA",
        address_code="3PL",
        ctns="18",
    ) == "DPSECO260812093-CA-3PL-18"


def test_parse_shipment_rename_excel():
    data = _xlsx_bytes(
        [
            ["序号", "运单号", "国家", "仓库代码", "收件人邮编", "件数"],
            [1, "DPSECO260812093", "CA", "3PL", "V5X 3Y1", 18],
            [2, "DPSECO260812094", "US", "AVP1", "18202-9361", 10],
            [3, "", "US", "BJC1", "93308", 5],
        ]
    )
    result = parse_shipment_rename_excel(data)
    assert result["total"] == 2
    assert result["items"][0]["targetStem"] == "DPSECO260812093-CA-3PL-18"
    assert result["items"][0]["folderName"] == ""
    assert result["items"][1]["targetStem"] == "DPSECO260812094-US-AVP1-10"


def test_parse_falls_back_to_zipcode_when_no_warehouse():
    data = _xlsx_bytes(
        [
            ["运单号", "国家", "仓库代码", "收件人邮编", "件数"],
            ["DPSECO1", "CA", "", "V5X 3Y1", 18],
            ["DPSECO2", "US", "AVP1", "18202", 10],
            ["DPSECO3", "US", "", "", 5],
        ]
    )
    result = parse_shipment_rename_excel(data)
    assert result["total"] == 2
    assert result["items"][0]["targetStem"] == "DPSECO1-CA-V5X3Y1-18"
    assert result["items"][0]["addressCode"] == "V5X 3Y1"
    assert result["items"][1]["targetStem"] == "DPSECO2-US-AVP1-10"
    assert any("DPSECO3" in e for e in result["errors"])


def test_parse_zipcode_only_column():
    data = _xlsx_bytes(
        [
            ["运单号", "国家", "邮编", "件数"],
            ["DPSECO1", "US", "93308", 5],
        ]
    )
    result = parse_shipment_rename_excel(data)
    assert result["total"] == 1
    assert result["items"][0]["targetStem"] == "DPSECO1-US-93308-5"


def test_parse_with_folder_name():
    data = _xlsx_bytes(
        [
            ["运单号", "国家", "仓库代码", "件数", "文件夹名"],
            ["DPSECO1", "CA", "3PL", 18, "A组"],
            ["DPSECO2", "US", "AVP1", 10, "A组"],
            ["DPSECO3", "US", "LAX9", 5, "B组"],
            ["DPSECO4", "US", "BJC1", 2, ""],
        ]
    )
    result = parse_shipment_rename_excel(data)
    assert result["total"] == 4
    assert result["items"][0]["folderName"] == "A组"
    assert result["items"][1]["folderName"] == "A组"
    assert result["items"][2]["folderName"] == "B组"
    assert result["items"][3]["folderName"] == ""


def test_parse_rejects_missing_headers():
    data = _xlsx_bytes([["运单号", "件数"], ["DPSECO1", 1]])
    try:
        parse_shipment_rename_excel(data)
        raise AssertionError("expected ValueError")
    except ValueError as exc:
        assert "表头" in str(exc)


def test_parse_skips_cover_sheet():
    from openpyxl import Workbook

    wb = Workbook()
    cover = wb.active
    cover.title = "封面"
    cover.append(["入仓单"])
    ws = wb.create_sheet("运单信息")
    ws.append(["运单号", "国家", "仓库代码", "件数"])
    ws.append(["DPSECO1", "CA", "3PL", 18])
    buf = BytesIO()
    wb.save(buf)
    result = parse_shipment_rename_excel(buf.getvalue())
    assert result["total"] == 1
    assert result["items"][0]["targetStem"] == "DPSECO1-CA-3PL-18"


def test_parse_header_with_newline_suffix():
    data = _xlsx_bytes(
        [
            ["运单号", "国家", "仓库代码", "件数\n（箱）"],
            ["DPSECO1", "CA", "3PL", 18],
        ]
    )
    result = parse_shipment_rename_excel(data)
    assert result["total"] == 1
