"""Generate channel_seeds.py from 渠道清单.xlsx."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

EXCEL = Path(r"c:\Users\admin\Desktop\youzi_bak\渠道清单.xlsx")
OUT = Path(__file__).resolve().parents[1] / "db" / "channel_seeds.py"


def infer_country(cn: str, name: str) -> str:
    blob = f"{cn} {name}"
    rules = [
        ("加拿大", "加拿大"),
        ("Canada", "加拿大"),
        ("欧洲", "欧洲"),
        ("European", "欧洲"),
        ("英国", "英国"),
        ("UK ", "英国"),
        ("澳大利亚", "澳大利亚"),
        ("AU ", "澳大利亚"),
        ("日本", "日本"),
        ("Japan", "日本"),
        ("阿联酋", "阿联酋"),
        ("UAE", "阿联酋"),
        ("Emirates", "阿联酋"),
        ("墨西哥", "墨西哥"),
        ("Mexico", "墨西哥"),
        ("海卡", "美国"),
        ("海派", "美国"),
        ("空派", "美国"),
        ("Sea Truck", "美国"),
        ("Sea Express", "美国"),
        ("Air Express", "美国"),
    ]
    for key, value in rules:
        if key in blob:
            return value
    stripped = blob.strip()
    if stripped.startswith("UK"):
        return "英国"
    if stripped.startswith("AU"):
        return "澳大利亚"
    return ""


def infer_category(cn: str, name: str) -> str:
    blob = f"{cn} {name}"
    if any(x in blob for x in ("空运", "空派", "Air")):
        return "空运"
    if any(x in blob for x in ("卡航", "Road")):
        return "卡航"
    if any(x in blob for x in ("铁路", "Train")):
        return "铁路"
    if any(
        x in blob
        for x in (
            "快船",
            "快速",
            "特快",
            "限时",
            "Rapid",
            "Priority",
            "Guaranteed",
            "Fast sea",
            "Fast maston",
        )
    ):
        return "快船"
    return "普船"


def main() -> None:
    wb = load_workbook(EXCEL, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))[1:]
    seen: set[str] = set()
    ordered: list[tuple[str, str, str, str, str, str, str]] = []
    for row in reversed(rows):
        code = str(row[3] or "").strip()
        if not code or code in seen:
            continue
        seen.add(code)
        cn = str(row[1] or "").strip()
        name = str(row[2] or "").strip()
        zh = cn or code
        ordered.append(
            (
                code,
                zh,
                name or code,
                infer_country(cn, name),
                infer_category(cn, name),
                "",
                zh,
            )
        )
    ordered.reverse()

    lines = [
        '"""渠道码表初始数据：来自「渠道清单.xlsx」全量。',
        "",
        "(code, name_zh, name_en, country, category, note, dps_channel_name, sort_order)",
        "",
        "- code：表格 code",
        "- name_zh / dps_channel_name：表格 cnName",
        "- name_en：表格 name",
        "- 表格内重复 code 保留最后一条",
        '"""',
        "",
        "from __future__ import annotations",
        "",
        'CHANNEL_CATEGORIES: tuple[str, ...] = ("快船", "普船", "卡航", "铁路", "空运")',
        "",
        "CHANNEL_SEEDS: list[tuple[str, str, str, str, str, str, str, int]] = [",
    ]
    for i, (code, zh, en, country, cat, note, dps) in enumerate(ordered):
        so = (i + 1) * 10
        lines.append(
            f"    ({code!r}, {zh!r}, {en!r}, {country!r}, {cat!r}, {note!r}, {dps!r}, {so}),"
        )
    lines.append("]")
    lines.append("")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {len(ordered)} seeds -> {OUT}")


if __name__ == "__main__":
    main()
